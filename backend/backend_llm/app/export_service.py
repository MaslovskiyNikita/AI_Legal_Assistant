import json
import re
from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

class ExportService:

    @staticmethod
    def _parse_markdown_to_docx(paragraph, text: str):
        """Парсит жирный и курсив и добавляет их как run в параграф Word"""
        pattern = r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*)'
        parts = re.split(pattern, text)
        
        for part in parts:
            if not part:
                continue
            if part.startswith('***') and part.endswith('***'):
                run = paragraph.add_run(part[3:-3])
                run.bold = True
                run.italic = True
            elif part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            else:
                paragraph.add_run(part)

    @classmethod
    def _render_content_docx(cls, doc, text: str):
        """Обрабатывает блоки текста: заголовки, списки и обычные абзацы"""
        if not text:
            return
            
        lines = text.split('\n')
        for line in lines:
            clean_line = line.strip()
            if not clean_line:
                doc.add_paragraph()
                continue
            
            if clean_line.startswith('### '):
                doc.add_heading(clean_line.replace('### ', ''), level=3)
            elif clean_line.startswith('## '):
                doc.add_heading(clean_line.replace('## ', ''), level=2)
            elif clean_line.startswith('# '):
                doc.add_heading(clean_line.replace('# ', ''), level=1)
            
            elif clean_line.startswith('* ') or clean_line.startswith('- '):
                p = doc.add_paragraph(style='List Bullet')
                cls._parse_markdown_to_docx(p, clean_line[2:])
                
            elif re.match(r'^\d+\.\s', clean_line):
                p = doc.add_paragraph(style='List Number')
                content = re.sub(r'^\d+\.\s', '', clean_line)
                cls._parse_markdown_to_docx(p, content)
                
            else:
                p = doc.add_paragraph()
                cls._parse_markdown_to_docx(p, clean_line)

    @classmethod
    def generate_docx_report(cls, diff_blocks_json: str, analysis) -> BytesIO:
        doc = Document()
        
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("ОТЧЕТ ОБ АНАЛИЗЕ ИЗМЕНЕНИЙ")
        run.bold = True
        run.font.size = Pt(16)

        doc.add_heading("Анализ рисков", level=2)
        risk_para = doc.add_paragraph()
        run_risk = risk_para.add_run(f"Общий уровень риска: {analysis.overall_risk.value}")
        run_risk.bold = True
        if analysis.overall_risk.value == "RED":
            run_risk.font.color.rgb = RGBColor(255, 0, 0)
        
        cls._render_content_docx(doc, analysis.summary)
         
        if analysis.details:
            doc.add_heading("Детальный анализ", level=3)
            for detail in analysis.details:
                p = doc.add_paragraph()
                r = p.add_run(f"• {detail.title}: [{detail.risk.value}]")
                r.bold = True
                if detail.risk.value == "RED":
                    r.font.color.rgb = RGBColor(255, 0, 0)
                elif detail.risk.value == "YELLOW":
                    r.font.color.rgb = RGBColor(204, 153, 0)
                    
                cls._render_content_docx(doc, detail.explanation)
                if detail.violated_law:
                    doc.add_paragraph(f"Законодательство: {detail.violated_law}").italic = True

        doc.add_heading("Таблица изменений", level=2)
        try:
            blocks = json.loads(diff_blocks_json)
            changed_blocks = [b for b in blocks if b.get('change_type') != 'UNCHANGED']
            if changed_blocks:
                table = doc.add_table(rows=1, cols=3)
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Тип'
                hdr_cells[1].text = 'Было'
                hdr_cells[2].text = 'Стало'
                for b in changed_blocks:
                    row_cells = table.add_row().cells
                    row_cells[0].text = b.get('change_type', '')
                    old_b = b.get('old_block')
                    new_b = b.get('new_block')
                    row_cells[1].text = old_b.get('text', '') if old_b else '-'
                    row_cells[2].text = new_b.get('text', '') if new_b else '-'
            else:
                doc.add_paragraph("Значимых изменений не найдено.")
        except Exception as e:
            doc.add_paragraph(f"Ошибка формирования таблицы: {str(e)}")

        out = BytesIO()
        doc.save(out)
        out.seek(0)
        return out


    @staticmethod
    def _md_to_html(text: str) -> str:
        """Преобразует Markdown в простой HTML для fpdf2"""
        text = text.replace('***', '<b><i>').replace('***', '</i></b>')
        text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
        text = re.sub(r'\*(.*?)\*', r'<i>\1</i>', text)
        text = text.replace('\n', '<br>')
        return text

    @classmethod
    def generate_pdf_report(cls, diff_blocks_json: str, analysis) -> BytesIO:
        pdf = FPDF()
        pdf.add_page()
        
        try:
            pdf.add_font('DejaVu', '', 'DejaVuSans.ttf')
            pdf.add_font('DejaVu', 'B', 'DejaVuSans-Bold.ttf')
            pdf.add_font('DejaVu', 'I', 'DejaVuSans-Oblique.ttf')
            pdf.set_font('DejaVu', '', 12)
        except:
            pdf.set_font("Arial", size=12)

        pdf.set_font('DejaVu', 'B', 16)
        pdf.cell(0, 10, "ОТЧЕТ ОБ АНАЛИЗЕ ИЗМЕНЕНИЙ", ln=True, align='C')
        pdf.ln(5)

        pdf.set_font('DejaVu', 'B', 14)
        pdf.cell(0, 10, "Анализ рисков", ln=True)
        
        if analysis.overall_risk.value == "RED":
            pdf.set_text_color(255, 0, 0)
        pdf.set_font('DejaVu', 'B', 12)
        pdf.cell(0, 10, f"Общий уровень риска: {analysis.overall_risk.value}", ln=True)
        pdf.set_text_color(0, 0, 0)
        
        pdf.set_font('DejaVu', '', 11)
        pdf.write_html(cls._md_to_html(analysis.summary))
        pdf.ln(10)

        if analysis.details:
            pdf.set_font('DejaVu', 'B', 13)
            pdf.cell(0, 10, "Детальный анализ", ln=True)
            for detail in analysis.details:
                header_html = f"<b>• {detail.title} [{detail.risk.value}]</b>"
                pdf.write_html(header_html)
                pdf.ln(5)
                
                pdf.write_html(cls._md_to_html(detail.explanation))
                
                if detail.violated_law:
                    pdf.ln(2)
                    pdf.write_html(f"<i>Связано с: {detail.violated_law}</i>")
                pdf.ln(8)

        pdf.ln(5)
        pdf.set_font('DejaVu', 'B', 13)
        pdf.multi_cell(0, 10, "Таблица изменений")
        
        try:
            blocks = json.loads(diff_blocks_json)
            changed = [b for b in blocks if b.get('change_type') != 'UNCHANGED']
            if changed:
                pdf.set_font('DejaVu', 'B', 10)
                col_width = pdf.epw / 3
                
                pdf.cell(col_width, 10, "Тип", border=1)
                pdf.cell(col_width, 10, "Было", border=1)
                pdf.cell(col_width, 10, "Стало", border=1)
                pdf.ln()
                
                pdf.set_font('DejaVu', '', 9)
                for b in changed:
                    old_txt = b.get('old_block', {}).get('text', '-') if b.get('old_block') else '-'
                    new_txt = b.get('new_block', {}).get('text', '-') if b.get('new_block') else '-'
                    
                    max_lines = max(pdf.get_nb_lines(col_width, old_txt), pdf.get_nb_lines(col_width, new_txt))
                    line_height = 6
                    row_h = max_lines * line_height
                    
                    if pdf.get_y() + row_h > 270:
                        pdf.add_page()

                    x = pdf.get_x()
                    y = pdf.get_y()
                    
                    pdf.multi_cell(col_width, row_h / max(1, pdf.get_nb_lines(col_width, b.get('change_type',''))), b.get('change_type', ''), border=1)
                    pdf.set_xy(x + col_width, y)
                    pdf.multi_cell(col_width, line_height, old_txt, border=1)
                    pdf.set_xy(x + col_width * 2, y)
                    pdf.multi_cell(col_width, line_height, new_txt, border=1)
                    pdf.set_y(y + row_h)
            else:
                pdf.cell(0, 10, "Изменений не найдено", ln=True)
        except Exception as e:
            pdf.cell(0, 10, f"Ошибка формирования таблицы: {str(e)}", ln=True)

        out = BytesIO()
        pdf_bytes = pdf.output()
        out.write(pdf_bytes)
        out.seek(0)
        return out