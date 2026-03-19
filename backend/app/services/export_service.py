import io
import asyncio
import re
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF
from loguru import logger


def clean_xml_string(text: str) -> str:
    if not text:
        return ""
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)


def _md_to_html(text: str) -> str:
    """Преобразует Markdown в простой HTML для FPDF."""
    if not text:
        return ""
    # Жирный курсив
    text = re.sub(r'\*\*\*(.*?)\*\*\*', r'<b><i>\1</i></b>', text)
    # Жирный
    text = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', text)
    # Курсив
    text = re.sub(r'\*(.*?)\*', r'<i>\1</i>', text)
    # Переносы строк
    text = text.replace('\n', '<br>')
    return text


def _parse_markdown_to_docx(paragraph, text: str):
    """Парсит Markdown и добавляет отформатированные runs в DOCX."""
    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i > 0:
            paragraph.add_run('\n')  # Добавляем перенос строки внутри параграфа

        pattern = r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*)'
        parts = re.split(pattern, line)

        for part in parts:
            if not part:
                continue
            if part.startswith('***') and part.endswith('***'):
                run = paragraph.add_run(part[3:-3])
                run.bold = True
                run.italic = True
            elif part.startswith('**') and part.endswith('**'):
                run = paragraph.add_run(part[2:-2])
                run.bold = True
            elif part.startswith('*') and part.endswith('*'):
                run = paragraph.add_run(part[1:-1])
                run.italic = True
            else:
                run = paragraph.add_run(part)

            run.font.size = Pt(11)


class ChatPDF(FPDF):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.roboto_loaded = False

        reg_path = "/app/fonts/Roboto-Regular.ttf"
        med_path = "/app/fonts/Roboto-Medium.ttf"

        try:
            if os.path.exists(reg_path) and os.path.exists(med_path):
                self.add_font("Roboto", "", reg_path)
                self.add_font("Roboto", "B", med_path)
                self.roboto_loaded = True
                logger.info("✅ PDF: Шрифты Roboto успешно подключены")
            else:
                logger.error(f"❌ PDF: Шрифты не найдены по путям: {reg_path}, {med_path}")
        except Exception as e:
            logger.error(f"❌ PDF: Ошибка загрузки шрифтов: {e}")

    def header(self):
        font_name = "Roboto" if self.roboto_loaded else "Helvetica"

        self.set_font(font_name, "B", 16)
        self.set_text_color(51, 144, 236)
        self.cell(0, 10, "AI Legal Expert", align="L")

        self.set_y(10)
        self.set_font(font_name, "", 10)
        self.set_text_color(142, 142, 147)
        meta_text = f"Сгенерировано: {datetime.now().strftime('%d.%m.%Y')}" if self.roboto_loaded else f"Generated: {datetime.now().strftime('%d.%m.%Y')}"
        self.cell(0, 10, meta_text, align="R")
        self.ln(15)

        self.set_draw_color(229, 229, 234)
        self.line(15, self.get_y(), 195, self.get_y())
        self.ln(10)


def _build_pdf_sync(messages: list) -> io.BytesIO:
    pdf = ChatPDF()
    pdf.add_page()

    font_name = pdf.font_family
    pdf.set_font(font_name, "B", 14)
    pdf.set_text_color(0, 0, 0)

    title_text = "Отчет о сравнении документов" if pdf.roboto_loaded else "Chat History Report"
    pdf.cell(0, 10, title_text, align="C", ln=True)
    pdf.ln(10)

    for msg in messages:
        role = msg.get('role', 'user')
        text = clean_xml_string(msg.get('text', '') or msg.get('content', ''))

        pdf.set_font(font_name, "B", 11)
        if role == "ai":
            pdf.set_text_color(51, 144, 236)
            prefix = "Legal Expert AI"
        else:
            pdf.set_text_color(0, 0, 0)
            prefix = "Пользователь" if pdf.roboto_loaded else "User"

        pdf.cell(0, 8, f"{prefix}:", ln=True)

        pdf.set_font(font_name, "", 10)
        pdf.set_text_color(0, 0, 0)

        # Конвертируем Markdown в HTML и рендерим
        html_text = _md_to_html(text)
        pdf.write_html(html_text)
        pdf.ln(5)

    stream = io.BytesIO(pdf.output())
    stream.seek(0)
    return stream


def _build_docx_sync(messages: list) -> io.BytesIO:
    doc = Document()
    date_str = datetime.now().strftime("%d.%m.%Y")

    section = doc.sections[0]
    header = section.header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header_para.add_run(f"AI Legal Expert | Сгенерировано: {date_str}")
    run.font.color.rgb = RGBColor(142, 142, 147)
    run.font.size = Pt(10)

    title = doc.add_heading('Отчет о сравнении документов', level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for msg in messages:
        role = msg.get('role', 'user')
        text = clean_xml_string(msg.get('text', '') or msg.get('content', ''))
        prefix = "Пользователь" if role == "user" else "Legal Expert AI"

        p = doc.add_paragraph()
        name_run = p.add_run(f"{prefix}:\n")
        name_run.bold = True
        if role == "ai":
            name_run.font.color.rgb = RGBColor(51, 144, 236)

        # Пропускаем текст через наш парсер
        _parse_markdown_to_docx(p, text)

    footer = section.footer
    f_para = footer.paragraphs[0]
    f_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = f_para.add_run("Документ подготовлен с помощью нейросети. Рекомендуется финальная проверка юристом.")
    f_run.font.color.rgb = RGBColor(142, 142, 147)
    f_run.font.size = Pt(9)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    return stream


async def generate_docx_stream(messages: list) -> io.BytesIO:
    """Запускает генерацию DOCX в отдельном потоке, не блокируя FastAPI"""
    return await asyncio.to_thread(_build_docx_sync, messages)


async def generate_pdf_stream(messages: list) -> io.BytesIO:
    """Запускает генерацию PDF в отдельном потоке, не блокируя FastAPI"""
    return await asyncio.to_thread(_build_pdf_sync, messages)