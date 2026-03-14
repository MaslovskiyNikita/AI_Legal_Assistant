from io import BytesIO
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .models import BlockDiff, ChangeType

class ExportService:
    @staticmethod
    def generate_docx_report(diffs: list[BlockDiff]) -> BytesIO:
        doc = Document()
        
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("ОТЧЕТ ОБ АНАЛИЗЕ ИЗМЕНЕНИЙ")
        run.bold = True
        run.font.size = Pt(16)

        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        headers = ["№", "Статус", "Содержание изменения", "Анализ рисков"]
        for i, text in enumerate(headers):
            table.cell(0, i).text = text

        row_idx = 1
        for diff in diffs:
            if diff.change_type == ChangeType.UNCHANGED:
                continue
                
            row = table.add_row()
            row.cells[0].text = str(row_idx)
            row.cells[1].text = diff.change_type.value
            
            # Текст изменений
            cell_text = row.cells[2]
            if diff.old_block:
                p = cell_text.add_paragraph()
                p.add_run(f"БЫЛО: {diff.old_block.text}").italic = True
            if diff.new_block:
                p = cell_text.add_paragraph()
                p.add_run(f"СТАЛО: {diff.new_block.text}").bold = True

            # Риски
            cell_risk = row.cells[3]
            p_risk = cell_risk.add_paragraph()
            run_risk = p_risk.add_run(f"РИСК: {diff.risk or 'НЕ ОПРЕДЕЛЕН'}")
            run_risk.bold = True
            
            if diff.comment:
                cell_risk.add_paragraph(diff.comment)
            if diff.violated_law:
                cell_risk.add_paragraph(f"Нарушает: {diff.violated_law}").italic = True
                
            row_idx += 1

        out = BytesIO()
        doc.save(out)
        out.seek(0)
        return out